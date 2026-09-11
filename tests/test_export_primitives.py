"""Generic editor primitives must survive the exporters.

Tables arrive as GFM pipe tables and task lists as GFM ``- [ ]`` markers —
neither is raw HTML. python-markdown renders the table itself, and the small
treeprocessor in ``app.services.export`` turns task markers into checkboxes so
DOCX, PDF and EPUB do not print literal ``[ ]``.
"""
from __future__ import annotations

import io
import re
import zipfile
import zlib

from docx import Document

from app.services import documents as documents_service
from app.services import projects as projects_service
from app.services.export import md_to_html

BODY = (
    "| Name | Role |\n"
    "| --- | --- |\n"
    "| Alice | Lead |\n"
    "| Bob | Editor |\n"
    "\n"
    "- [ ] draft\n"
    "- [x] outline"
)

MARKS = (
    'plain <mark>hi</mark> a<sub>2</sub> b<sup>x</sup> '
    '<span style="color:#c00">red</span> end'
)


def _project_with(make_project, body: str) -> None:
    make_project("p1")
    documents_service.create_document("p1", "Chapter One", "chapter", content=body)


def _pdf_text(data: bytes) -> bytes:
    text = b""
    for match in re.finditer(rb"stream\r?\n", data):
        start = match.end()
        end = data.find(b"endstream", start)
        chunk = data[start:end]
        try:
            text += zlib.decompress(chunk)
        except zlib.error:
            text += chunk
    return text


def _epub_html(data: bytes) -> str:
    archive = zipfile.ZipFile(io.BytesIO(data))
    return "\n".join(
        archive.read(name).decode("utf-8", "replace")
        for name in archive.namelist()
        if name.endswith(".xhtml")
    )


def test_md_to_html_renders_tables_and_task_glyphs():
    html = md_to_html(BODY)
    assert "<table>" in html
    assert "<th>Name</th>" in html
    assert "\u2610 draft" in html, "an unchecked marker becomes a checkbox glyph"
    assert "\u2611 outline" in html, "a checked marker becomes a checked glyph"
    assert "[ ] draft" not in html
    assert "[x] outline" not in html


def test_docx_keeps_a_gfm_table_and_task_glyphs(make_project):
    _project_with(make_project, BODY)

    doc = Document(io.BytesIO(projects_service.export_docx("p1")))
    assert len(doc.tables) == 1
    assert [cell.text.strip() for cell in doc.tables[0].rows[0].cells] == ["Name", "Role"]
    assert [cell.text.strip() for cell in doc.tables[0].rows[1].cells] == ["Alice", "Lead"]

    paragraph_text = "\n".join(p.text for p in doc.paragraphs)
    assert "\u2610 draft" in paragraph_text
    assert "\u2611 outline" in paragraph_text


def test_pdf_keeps_the_table_and_task_text(make_project, monkeypatch):
    _project_with(make_project, BODY)
    # Force the built-in core fonts; the glyphs degrade to "?" but the words stay.
    monkeypatch.setattr(projects_service, "_find_pdf_fonts", lambda: {"serif": None, "mono": None})

    data = projects_service.export_pdf("p1")
    assert data[:5] == b"%PDF-"
    text = _pdf_text(data)
    assert b"Alice" in text
    assert b"Editor" in text
    assert b"draft" in text


def test_epub_keeps_the_primitive_markup(make_project):
    _project_with(make_project, BODY)

    html = _epub_html(projects_service.export_epub("p1"))
    assert "<table>" in html
    assert "<th>Name</th>" in html
    assert "\u2610 draft" in html
    assert "\u2611 outline" in html


def test_epub_keeps_inline_mark_markup(make_project):
    _project_with(make_project, MARKS)

    html = _epub_html(projects_service.export_epub("p1"))
    assert "<mark>hi</mark>" in html
    assert "<sub>2</sub>" in html
    assert "<sup>x</sup>" in html
    assert "color:#c00" in html


def test_docx_keeps_inline_mark_text(make_project):
    _project_with(make_project, MARKS)

    doc = Document(io.BytesIO(projects_service.export_docx("p1")))
    text = "\n".join(p.text for p in doc.paragraphs)
    # The styling is not carried into Word, but no text may be lost.
    for word in ("hi", "2", "x", "red"):
        assert word in text
    assert "[ ]" not in text
