"""Import route tests: the multipart shape the browser actually sends.

The client posts every picked file as its own part plus a ``paths`` array
carrying the relative paths in the same order (multipart alone loses a folder
pick's structure). These tests pin that contract, the status codes for the bad
cases, and the localhost guard.
"""
from __future__ import annotations

import io
import json
import zipfile

import pytest
from fastapi.testclient import TestClient

from app.main import create_app
from app.services import documents as documents_service

HOST = {"host": "127.0.0.1"}


@pytest.fixture
def api_data_dir(tmp_path):
    """A throwaway data dir owned by the API tests (create_app makes it itself)."""
    d = tmp_path / "data"
    d.mkdir(parents=True, exist_ok=True)
    return d


@pytest.fixture
def client(api_data_dir, monkeypatch):
    monkeypatch.setenv("IWAKURA_DATA_DIR", str(api_data_dir))
    with TestClient(create_app()) as test_client:
        yield test_client


def make_project(client, project_id: str = "p1") -> dict:
    resp = client.post("/api/projects", json={"name": project_id}, headers=HOST)
    assert resp.status_code in (200, 201), resp.text
    return resp.json()


def post_import(client, project_id, parts, *, paths=None, **fields):
    """Post multipart the way static/js/api.js does."""
    files = [("files", (name, data, "application/octet-stream")) for name, data in parts]
    data = {
        "paths": json.dumps(paths if paths is not None else [name for name, _ in parts]),
        "folder": "",
        "name": "",
        "as": "chapter",
        "source": "auto",
    }
    data.update(fields)
    return client.post(f"/api/projects/{project_id}/import", files=files, data=data, headers=HOST)


def test_import_uploads_a_folder_of_markdown(client, api_data_dir):
    make_project(client)
    resp = post_import(
        client,
        "p1",
        [
            ("01-open.md", b"---\ntitle: The Opening\n---\nIt began."),
            ("02-two.md", b"Second."),
        ],
        paths=["Novel/01-open.md", "Novel/02-two.md"],
        name="Novel",
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["documents"] == 2
    assert body["isFolder"] is True
    assert body["folder"].endswith("Novel")

    ids = [doc["id"] for doc in documents_service.iter_documents("p1")]
    assert ids == [f"{body['folder']}/01-the-opening", f"{body['folder']}/02-two"]


def test_import_accepts_a_zip_bundle(client):
    make_project(client)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("Novel/one.md", "One")
        archive.writestr("Novel/two.md", "Two")
    resp = post_import(client, "p1", [("novel.zip", buf.getvalue())], paths=["novel.zip"])
    assert resp.status_code == 201, resp.text
    assert resp.json()["documents"] == 2


def test_import_into_a_named_folder(client):
    make_project(client)
    folder = client.post("/api/projects/p1/folders", json={"name": "Act One"}, headers=HOST).json()
    resp = post_import(
        client,
        "p1",
        [("a.md", b"a"), ("b.md", b"b")],
        folder=folder["id"],
    )
    assert resp.status_code == 201, resp.text
    assert all(doc["id"].startswith(folder["id"] + "/") for doc in documents_service.iter_documents("p1"))


def test_no_files_is_rejected(client):
    make_project(client)
    resp = client.post("/api/projects/p1/import", data={"paths": "[]"}, headers=HOST)
    assert resp.status_code == 422


def test_unreadable_files_are_a_400(client):
    make_project(client)
    resp = post_import(client, "p1", [("scan.pdf", b"%PDF-1.4 junk")])
    assert resp.status_code == 400
    assert "imported" in resp.json()["detail"].lower()


def test_a_corrupt_word_document_is_a_400(client):
    make_project(client)
    resp = post_import(client, "p1", [("notes.docx", b"PK\x03\x04not really a package")])
    assert resp.status_code == 400
    assert "word" in resp.json()["detail"].lower()


def test_empty_text_files_are_a_400(client):
    make_project(client)
    resp = post_import(client, "p1", [("empty.bin", b"x")])
    assert resp.status_code == 400


def test_a_missing_project_is_a_404(client):
    resp = post_import(client, "nope", [("a.md", b"hi")])
    assert resp.status_code == 404


def test_a_foreign_origin_is_refused(client):
    make_project(client)
    resp = client.post(
        "/api/projects/p1/import",
        files=[("files", ("a.md", b"hi", "text/markdown"))],
        data={"paths": '["a.md"]'},
        headers={"host": "127.0.0.1", "origin": "http://evil.example"},
    )
    assert resp.status_code == 403


def make_docx() -> bytes:
    """A two-chapter Word document, built with python-docx."""
    from docx import Document

    document = Document()
    document.add_heading("Chapter One", level=1)
    document.add_paragraph("First.")
    document.add_heading("Chapter Two", level=1)
    document.add_paragraph("Second.")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_import_reads_a_word_document(client):
    make_project(client)
    resp = post_import(client, "p1", [("Manuscript.docx", make_docx())])
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["documents"] == 2
    assert body["source"] == "docx"
    titles = [doc["title"] for doc in documents_service.iter_documents("p1")]
    assert titles == ["Chapter One", "Chapter Two"]


def test_the_split_field_turns_word_splitting_off(client):
    make_project(client)
    resp = post_import(client, "p1", [("Manuscript.docx", make_docx())], split="0")
    assert resp.status_code == 201, resp.text
    assert resp.json()["documents"] == 1
    docs = documents_service.iter_documents("p1")
    assert docs[0]["title"] == "Manuscript"
    assert "# Chapter One" in docs[0]["body"]


def test_import_as_notes_is_honoured(client):
    make_project(client)
    resp = post_import(client, "p1", [("a.md", b"a"), ("b.md", b"b")], **{"as": "note"})
    assert resp.status_code == 201, resp.text
    kinds = {doc["kind"] for doc in documents_service.iter_documents("p1")}
    assert kinds == {"note"}


def test_paths_fall_back_to_the_upload_filename(client):
    make_project(client)
    resp = post_import(client, "p1", [("Chapter One.md", b"hi")], paths=[])
    assert resp.status_code == 201, resp.text
    assert resp.json()["folder"] == "01-chapter-one"
