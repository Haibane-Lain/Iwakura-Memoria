"""Exports must carry vertical timelines.

A timeline is raw HTML in the Markdown (``<aside class="timeline">`` with a
nested ``<table>``), so every exporter meets it: DOCX needs real Word-table
support, PDF renders the nested table through fpdf2, and EPUB keeps the markup
and its own stylesheet. A malformed or empty timeline must never take an export
down.
"""
from __future__ import annotations

import io
import re
import zipfile

from docx import Document

from app import config
from app.services import documents as documents_service
from app.services import projects as projects_service


def _project_with(make_project, body: str = "") -> None:
    folder = make_project("p1")
    (folder / config.WIKI_DIRNAME).mkdir(exist_ok=True)
    documents_service.create_document("p1", "Chapter One", "chapter", content=body)


def _timeline() -> str:
    return "\n".join(
        [
            '<aside class="timeline">',
            '<table class="tl-rows">',
            '<tr class="tl-title"><th colspan="2">Timeline of the War</th></tr>',
            '<tr class="tl-section"><th colspan="2">Early Era</th></tr>',
            '<tr class="tl-event"><td class="tl-date">1066</td>'
            '<td class="tl-body">Norman Conquest</td></tr>',
            '<tr class="tl-event"><td class="tl-date">1215</td>'
            '<td class="tl-body">Magna Carta</td></tr>',
            "</table>",
            "</aside>",
        ]
    )


def _docx(data: bytes) -> tuple[Document, zipfile.ZipFile]:
    archive = zipfile.ZipFile(io.BytesIO(data))
    return Document(io.BytesIO(data)), archive


def _table_text(doc: Document) -> list[list[str]]:
    return [
        [cell.text.strip() for cell in row.cells]
        for table in doc.tables
        for row in table.rows
    ]


def _pdf_text(data: bytes) -> bytes:
    """The text of a PDF, with fpdf2's compressed content streams inflated.

    Only meaningful when the export uses fpdf2's built-in core fonts: a Unicode
    TTF is embedded as a subset with its own encoding, so the words are no
    longer literal bytes in the stream (see ``_core_fonts``).
    """
    import zlib

    text = b""
    for match in re.finditer(rb"stream\r?\n", data):
        start = match.end()
        end = data.find(b"endstream", start)
        chunk = data[start:end]
        try:
            text += zlib.decompress(chunk)
        except zlib.error:
            text += chunk
    return text


def _core_fonts(monkeypatch) -> None:
    """Force fpdf2's built-in fonts so the PDF's words can be asserted."""
    monkeypatch.setattr(
        projects_service, "_find_pdf_fonts", lambda: {"serif": None, "mono": None}
    )


def _epub_parts(data: bytes) -> tuple[zipfile.ZipFile, str]:
    archive = zipfile.ZipFile(io.BytesIO(data))
    html = "\n".join(
        archive.read(name).decode("utf-8", "replace")
        for name in archive.namelist()
        if name.endswith(".xhtml")
    )
    return archive, html


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def test_docx_keeps_the_timeline_as_a_real_table(make_project):
    _project_with(make_project, _timeline())

    doc, _ = _docx(projects_service.export_docx("p1"))
    assert len(doc.tables) == 1
    rows = _table_text(doc)
    assert [row[0] for row in rows] == [
        "Timeline of the War",
        "Early Era",
        "1066",
        "1215",
    ]
    # Each event keeps its date and description as two cells.
    assert rows[2][1] == "Norman Conquest"
    assert rows[3][1] == "Magna Carta"


def test_docx_survives_a_malformed_timeline(make_project):
    _project_with(make_project)
    body = (
        '<aside class="timeline">\n<table class="tl-rows">\n'
        '<tr class="tl-title"><th colspan="2">Only a title</th></tr>\n'
        '<tr class="tl-event"><td class="tl-date">1066</td></tr>\n'
        "</table>\n</aside>"
    )
    documents_service.save_document("p1", "01-chapter-one", body)

    doc, _ = _docx(projects_service.export_docx("p1"))
    assert len(doc.tables) == 1
    assert _table_text(doc)[1][0] == "1066"


def test_docx_handles_a_document_without_a_timeline(make_project):
    _project_with(make_project, "Just prose.\n\n- and a list")

    doc, _ = _docx(projects_service.export_docx("p1"))
    assert doc.tables == []
    assert any("Just prose." in p.text for p in doc.paragraphs)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_pdf_renders_the_timeline_table(make_project, monkeypatch):
    _project_with(make_project, _timeline())
    _core_fonts(monkeypatch)

    data = projects_service.export_pdf("p1")
    assert data[:5] == b"%PDF-"
    text = _pdf_text(data)
    assert b"Timeline of the War" in text
    assert b"Early Era" in text
    assert b"1066" in text
    assert b"Norman Conquest" in text
    # Two columns: the date and the description sit at different x offsets on
    # the same line of the table.
    offsets = [float(value) for value in re.findall(rb"([\d.]+) [\d.]+ Td", text)]
    assert len(set(round(value) for value in offsets)) > 1


def test_pdf_survives_a_timeline_with_a_missing_body(make_project, monkeypatch):
    _project_with(make_project, _timeline().replace("Norman Conquest", ""))
    _core_fonts(monkeypatch)

    data = projects_service.export_pdf("p1")
    assert data[:5] == b"%PDF-"
    assert b"1066" in _pdf_text(data)


# ---------------------------------------------------------------------------
# EPUB
# ---------------------------------------------------------------------------


def test_epub_keeps_the_timeline_markup_and_styles_it(make_project):
    _project_with(make_project, _timeline())

    archive, html = _epub_parts(projects_service.export_epub("p1"))
    assert '<aside class="timeline"' in html
    assert '<table class="tl-rows">' in html
    assert '<td class="tl-date">1066</td>' in html
    styles = [
        archive.read(name).decode("utf-8", "replace")
        for name in archive.namelist()
        if name.endswith(".css")
    ]
    assert any("aside.timeline" in sheet for sheet in styles), styles
