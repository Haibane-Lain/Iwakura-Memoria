"""Exports must carry inline pictures (and never break on a bad one).

DOCX / PDF / EPUB all have to resolve a document's ``assets/<name>`` reference
to a real file and embed it. A picture that cannot be resolved — a remote URL,
a hand-deleted file, a corrupt blob — must be *dropped*, not passed to a
converter: fpdf2 raises on a source it cannot open, which used to be able to
take the whole PDF export down with it.
"""
from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.shared import Inches
from PIL import Image

from app import config
from app.services import assets as assets_service
from app.services import documents as documents_service
from app.services import projects as projects_service


def png_bytes(width: int = 8, height: int = 6) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


def shape_width_inches(doc) -> float:
    return doc.inline_shapes[0].width / 914400


def _project_with(make_project, body: str = ""):
    folder = make_project("p1")
    (folder / config.WIKI_DIRNAME).mkdir(exist_ok=True)
    documents_service.create_document("p1", "Chapter One", "chapter", content=body)
    return folder


def _stored_image(name: str = "mara.png") -> str:
    return assets_service.save_image("p1", name, png_bytes())["path"]


def _epub_parts(data: bytes) -> tuple[zipfile.ZipFile, str]:
    archive = zipfile.ZipFile(io.BytesIO(data))
    html = "\n".join(
        archive.read(name).decode("utf-8", "replace")
        for name in archive.namelist()
        if name.endswith(".xhtml")
    )
    return archive, html


def _epub_images(archive: zipfile.ZipFile) -> list[str]:
    """Image resources inside the epub, as the chapter HTML names them.

    EbookLib stores everything under ``EPUB/``, while the ``src`` it is
    referenced by keeps the plain ``images/…`` file name.
    """
    out = []
    for name in archive.namelist():
        inner = name[len("EPUB/"):] if name.startswith("EPUB/") else name
        if inner.startswith("images/"):
            out.append(inner)
    return out


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def test_docx_embeds_a_plain_markdown_image(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document("p1", "01-chapter-one", f"Before\n\n![]({path})\n\nAfter")

    doc = Document(io.BytesIO(projects_service.export_docx("p1")))
    assert len(doc.inline_shapes) == 1
    assert doc.inline_shapes[0].width <= Inches(6.5)


def test_docx_embeds_a_resized_html_image(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document(
        "p1", "01-chapter-one", f'text <img src="{path}" alt="Mara" width="120"> more'
    )

    doc = Document(io.BytesIO(projects_service.export_docx("p1")))
    assert len(doc.inline_shapes) == 1
    # 120 CSS px at 96 dpi = 1.25 in.
    assert abs(shape_width_inches(doc) - 1.25) < 0.02


def test_docx_drops_a_missing_image(make_project):
    _project_with(make_project)
    documents_service.save_document("p1", "01-chapter-one", "text ![](assets/nope.png) more")

    doc = Document(io.BytesIO(projects_service.export_docx("p1")))
    assert len(doc.inline_shapes) == 0


def test_docx_ignores_remote_images(make_project):
    _project_with(make_project)
    documents_service.save_document(
        "p1", "01-chapter-one", "text ![](https://example.com/tracker.png) more"
    )

    doc = Document(io.BytesIO(projects_service.export_docx("p1")))
    assert len(doc.inline_shapes) == 0


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_pdf_embeds_an_image(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document("p1", "01-chapter-one", f"Before\n\n![]({path})\n\nAfter")

    data = projects_service.export_pdf("p1")
    assert data[:5] == b"%PDF-"
    assert b"/Subtype /Image" in data


def test_pdf_survives_a_missing_image(make_project):
    _project_with(make_project)
    documents_service.save_document("p1", "01-chapter-one", "text ![](assets/gone.png) more")

    data = projects_service.export_pdf("p1")
    assert data[:5] == b"%PDF-"
    assert b"/Subtype /Image" not in data


def test_pdf_survives_a_resized_image(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document(
        "p1", "01-chapter-one", f'text <img src="{path}" width="90"> more'
    )
    assert projects_service.export_pdf("p1")[:5] == b"%PDF-"


# ---------------------------------------------------------------------------
# EPUB
# ---------------------------------------------------------------------------


def test_epub_embeds_an_image_and_points_at_it(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document("p1", "01-chapter-one", f"Before\n\n![]({path})\n\nAfter")

    archive, html = _epub_parts(projects_service.export_epub("p1"))
    images = _epub_images(archive)
    assert len(images) == 1
    assert images[0].endswith(".png")
    assert f'src="{images[0]}"' in html
    assert "assets/" not in html


def test_epub_keeps_a_resized_width(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document(
        "p1", "01-chapter-one", f'text <img src="{path}" width="300"> more'
    )

    archive, html = _epub_parts(projects_service.export_epub("p1"))
    assert len(_epub_images(archive)) == 1
    assert 'width="300"' in html


def test_epub_drops_a_missing_image(make_project):
    _project_with(make_project)
    documents_service.save_document("p1", "01-chapter-one", "text ![](assets/gone.png) more")

    archive, html = _epub_parts(projects_service.export_epub("p1"))
    assert _epub_images(archive) == []
    assert "<img" not in html


def test_epub_shares_one_resource_for_a_repeated_image(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document(
        "p1", "01-chapter-one", f"![]({path})\n\nagain:\n\n![]({path})"
    )

    archive, _ = _epub_parts(projects_service.export_epub("p1"))
    assert len(_epub_images(archive)) == 1


# ---------------------------------------------------------------------------
# ZIP
# ---------------------------------------------------------------------------


def test_zip_includes_pictures_even_for_a_filtered_export(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document("p1", "01-chapter-one", f"![]({path})")

    archive = zipfile.ZipFile(io.BytesIO(projects_service.export_zip("p1", ["."])))
    names = archive.namelist()
    assert "chapter-one.md" in names
    assert any(name.endswith(path.split("/")[-1]) for name in names), names
