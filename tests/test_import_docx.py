"""DOCX import tests: a Word document, read as import-outline documents.

Fixtures are built with python-docx in-test, so every construct under test is a
real Word part rather than a hand-written approximation of one. The reader walks
the body XML (tables and pictures keep their place, hyperlinks are readable at
all), so these tests pin the Markdown it produces, the split behaviour, and the
things it deliberately refuses.
"""
from __future__ import annotations

import io

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from app import config
from app.services import documents as documents_service
from app.services import import_docs as import_service


def png_bytes(width: int = 8, height: int = 6) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


def build_docx(fill) -> bytes:
    document = Document()
    fill(document)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def add_hyperlink(paragraph, url: str, text: str) -> None:
    """Word hyperlinks have no python-docx API, so the XML is built by hand."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def import_one(make_project, raw: bytes, name: str = "Manuscript.docx", **options):
    """Import a single .docx into a fresh project; return its documents."""
    make_project("p1")
    summary = import_service.import_bundle("p1", None, [(name, raw)], options)
    return summary, documents_service.iter_documents("p1")


# ---------------------------------------------------------------------------
# structure
# ---------------------------------------------------------------------------


def test_headings_and_paragraphs_become_markdown(make_project):
    def fill(doc):
        doc.add_heading("Chapter One", level=1)
        doc.add_paragraph("The first line.")
        doc.add_heading("A Section", level=2)
        doc.add_paragraph("Deeper text.")

    summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert summary["documents"] == 1
    assert docs[0]["body"] == (
        "# Chapter One\n\nThe first line.\n\n## A Section\n\nDeeper text."
    )


def test_a_word_document_splits_at_heading_one(make_project):
    def fill(doc):
        doc.add_heading("Chapter One", level=1)
        doc.add_paragraph("First.")
        doc.add_heading("Chapter Two", level=1)
        doc.add_paragraph("Second.")
        doc.add_heading("Chapter Three", level=1)
        doc.add_paragraph("Third.")

    _summary, docs = import_one(make_project, build_docx(fill))
    assert [doc["title"] for doc in docs] == ["Chapter One", "Chapter Two", "Chapter Three"]
    assert [doc["body"] for doc in docs] == ["First.", "Second.", "Third."]


def test_content_before_the_first_heading_keeps_its_own_document(make_project):
    def fill(doc):
        doc.add_paragraph("A dedication.")
        doc.add_heading("Chapter One", level=1)
        doc.add_paragraph("First.")

    _summary, docs = import_one(make_project, build_docx(fill), name="03-the-book.docx")
    assert [doc["title"] for doc in docs] == ["The Book", "Chapter One"]
    assert docs[0]["body"] == "A dedication."


def test_a_document_without_headings_keeps_its_file_name(make_project):
    def fill(doc):
        doc.add_paragraph("Just prose.")

    _summary, docs = import_one(make_project, build_docx(fill), name="A Loose Note.docx")
    assert [doc["title"] for doc in docs] == ["A Loose Note"]
    assert docs[0]["body"] == "Just prose."


def test_headings_deeper_than_the_editor_renders_are_clamped(make_project):
    def fill(doc):
        doc.add_heading("One", level=1)
        doc.add_heading("Four", level=4)
        doc.add_heading("Six", level=6)

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert "### Four" in docs[0]["body"]
    assert "### Six" in docs[0]["body"]
    assert "####" not in docs[0]["body"]


def test_the_title_property_names_the_document(make_project):
    def fill(doc):
        doc.core_properties.title = "The Real Title"
        doc.add_paragraph("Body.")

    _summary, docs = import_one(make_project, build_docx(fill), name="whatever.docx", split=False)
    assert docs[0]["title"] == "The Real Title"


def test_a_legal_or_quote_paragraph_is_escaped(make_project):
    def fill(doc):
        doc.add_paragraph("#1 bestseller")  # not a heading: no space after #
        doc.add_paragraph("- not a list item")
        doc.add_paragraph("1. not a list either")

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert docs[0]["body"] == "#1 bestseller\n\n\\- not a list item\n\n1\\. not a list either"


def test_a_page_break_becomes_a_divider(make_project):
    def fill(doc):
        doc.add_paragraph("Before.")
        doc.add_page_break()
        doc.add_paragraph("After.")

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert docs[0]["body"] == "Before.\n\n---\n\nAfter."


def test_a_blank_document_still_imports(make_project):
    _summary, docs = import_one(make_project, build_docx(lambda doc: None), name="Empty.docx")
    assert len(docs) == 1
    assert docs[0]["body"] == ""


# ---------------------------------------------------------------------------
# inline formatting
# ---------------------------------------------------------------------------


def test_run_formatting_becomes_markdown(make_project):
    def fill(doc):
        para = doc.add_paragraph()
        para.add_run("bold").bold = True
        para.add_run(" ")
        para.add_run("italic").italic = True
        para.add_run(" ")
        both = para.add_run("both")
        both.bold = True
        both.italic = True
        para.add_run(" ")
        para.add_run("gone").font.strike = True
        para.add_run(" ")
        para.add_run("under").underline = True
        para.add_run(" ")
        sup = para.add_run("sup")
        sup.font.superscript = True
        para.add_run(" ")
        sub = para.add_run("sub")
        sub.font.subscript = True
        para.add_run(" ")
        code = para.add_run("code")
        code.font.name = "Consolas"
        para.add_run(" ")
        mark = para.add_run("lit")
        mark.font.highlight_color = 7  # WD_COLOR_INDEX.YELLOW

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    body = docs[0]["body"]
    assert "**bold**" in body
    assert "*italic*" in body
    assert "***both***" in body
    assert "~~gone~~" in body
    assert "<u>under</u>" in body
    assert "<sup>sup</sup>" in body
    assert "<sub>sub</sub>" in body
    assert "`code`" in body
    assert "<mark>lit</mark>" in body


def test_a_space_after_a_bold_run_stays_outside_the_markers(make_project):
    def fill(doc):
        para = doc.add_paragraph()
        para.add_run("Shouted ").bold = True
        para.add_run("quietly.")

    _summary, docs = import_one(make_project, docx := build_docx(fill), split=False)
    assert docs[0]["body"] == "**Shouted** quietly."


def test_a_hyperlink_becomes_a_markdown_link(make_project):
    def fill(doc):
        para = doc.add_paragraph("See ")
        add_hyperlink(para, "https://example.com/a?b=1", "the site")
        para.add_run(" now.")

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert docs[0]["body"] == "See [the site](https://example.com/a?b=1) now."


def test_tracked_insertions_are_kept_and_deletions_dropped(make_project):
    def fill(doc):
        para = doc.add_paragraph("Kept ")
        ins = OxmlElement("w:ins")
        run = OxmlElement("w:r")
        node = OxmlElement("w:t")
        node.text = "added"
        run.append(node)
        ins.append(run)
        para._p.append(ins)
        dele = OxmlElement("w:del")
        dead = OxmlElement("w:delText")
        dead.text = "removed"
        dead_run = OxmlElement("w:r")
        dead_run.append(dead)
        dele.append(dead_run)
        para._p.append(dele)

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert docs[0]["body"] == "Kept addedremoved".replace("removed", "")


# ---------------------------------------------------------------------------
# lists
# ---------------------------------------------------------------------------


def test_bulleted_and_numbered_lists_become_list_markdown(make_project):
    def fill(doc):
        doc.add_paragraph("bullet one", style="List Bullet")
        doc.add_paragraph("bullet two", style="List Bullet")
        doc.add_paragraph("first", style="List Number")
        doc.add_paragraph("second", style="List Number")
        doc.add_paragraph("after the lists")

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert docs[0]["body"] == (
        "- bullet one\n\n- bullet two\n\n1. first\n\n2. second\n\nafter the lists"
    )


def test_numbering_is_read_from_the_document_not_the_style_name(make_project):
    """A paragraph's style name need not say "number" — its numbering does.

    Word puts the level format in ``numbering.xml``; a "List Paragraph" using a
    decimal definition is an ordered list, and a bullet definition a bulleted
    one, even though neither style name hints at it.
    """

    def fill(doc):
        decimal = _add_numbering(doc, "decimal", "90")
        bullet = _add_numbering(doc, "bullet", "91")
        _numbered(doc, "one", decimal, 0)
        _numbered(doc, "two", decimal, 0)
        _numbered(doc, "an aside", bullet, 1)
        _numbered(doc, "back to the list", decimal, 0)

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert docs[0]["body"] == (
        "1. one\n\n2. two\n\n  - an aside\n\n3. back to the list"
    )


def _add_numbering(doc, fmt: str, abstract_id: str) -> str:
    """Register a numbering definition and return the numId paragraphs should use."""
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), abstract_id)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    level.append(num_fmt)
    abstract.append(level)
    numbering.insert(0, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), abstract_id)
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract_id)
    num.append(ref)
    numbering.append(num)
    return abstract_id


def _numbered(doc, text: str, num_id: str, level: int) -> None:
    """A plain-styled paragraph carrying an explicit numbering reference, as Word writes."""
    paragraph = doc.add_paragraph(text, style="List Paragraph")
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_pr.append(ilvl)
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), num_id)
    num_pr.append(num_ref)
    paragraph._p.get_or_add_pPr().append(num_pr)


# ---------------------------------------------------------------------------
# tables and pictures
# ---------------------------------------------------------------------------


def test_a_table_becomes_a_gfm_table(make_project):
    def fill(doc):
        table = doc.add_table(rows=3, cols=3)
        table.style = "Table Grid"
        values = [["Name", "Role", "Notes"], ["Mara", "Captain", "wary"], ["Ito", "Pilot", "cheerful"]]
        for r, row in enumerate(values):
            for c, value in enumerate(row):
                table.cell(r, c).text = value

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert docs[0]["body"] == (
        "| Name | Role | Notes |\n"
        "| --- | --- | --- |\n"
        "| Mara | Captain | wary |\n"
        "| Ito | Pilot | cheerful |"
    )


def test_a_pipe_inside_a_cell_cannot_break_the_row(make_project):
    def fill(doc):
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "a|b"
        table.cell(0, 1).text = "c"
        table.cell(1, 0).text = "d"
        table.cell(1, 1).text = "e"

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert "| a\\|b | c |" in docs[0]["body"]


def test_a_picture_lands_in_the_asset_folder(make_project):
    def fill(doc):
        doc.add_paragraph("Before the picture.")
        doc.add_picture(io.BytesIO(png_bytes()))

    summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert summary["images"] == 1
    body = docs[0]["body"]
    assert "@@" not in body, "the token must be replaced by the stored asset path"
    assert f"{config.ASSETS_DIRNAME}/" in body
    name = body.split(f"{config.ASSETS_DIRNAME}/")[1].split(")")[0]
    project = config.DATA_DIR / "p1"
    assert (project / config.ASSETS_DIRNAME / name).is_file()


def test_a_pictures_description_becomes_its_alt_text(make_project):
    def fill(doc):
        doc.add_picture(io.BytesIO(png_bytes()))
        # Word stores alt text on the drawing's docPr element.
        docpr = doc.element.body.find(".//" + qn("wp:docPr"))
        docpr.set("descr", "Mara's portrait")

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert "![Mara's portrait](" in docs[0]["body"]


def test_a_picture_only_appears_in_the_document_that_holds_it(make_project):
    def fill(doc):
        doc.add_heading("One", level=1)
        doc.add_paragraph("Text.")
        doc.add_heading("Two", level=1)
        doc.add_picture(io.BytesIO(png_bytes()))

    summary, docs = import_one(make_project, build_docx(fill))
    assert summary["images"] == 1
    assert "@@" not in docs[0]["body"], "chapter one must not carry chapter two's picture"
    assert "@@" not in docs[1]["body"]
    assert f"{config.ASSETS_DIRNAME}/" in docs[1]["body"]


# ---------------------------------------------------------------------------
# options and failures
# ---------------------------------------------------------------------------


def test_import_as_note_applies_to_word_documents(make_project):
    def fill(doc):
        doc.add_paragraph("prose")

    _summary, docs = import_one(make_project, build_docx(fill), **{"as": "note"})
    assert {doc["kind"] for doc in docs} == {"note"}


def test_splitting_can_be_turned_off(make_project):
    def fill(doc):
        doc.add_heading("One", level=1)
        doc.add_paragraph("a")
        doc.add_heading("Two", level=1)
        doc.add_paragraph("b")

    _summary, docs = import_one(make_project, build_docx(fill), split=False)
    assert len(docs) == 1
    assert "# One" in docs[0]["body"] and "# Two" in docs[0]["body"]


def test_a_missing_heading_1_makes_splitting_a_no_op(make_project):
    def fill(doc):
        doc.add_heading("Section", level=2)
        doc.add_paragraph("prose")

    _summary, docs = import_one(make_project, build_docx(fill), name="Notes.docx")
    assert len(docs) == 1
    assert docs[0]["title"] == "Notes"
    assert docs[0]["body"] == "## Section\n\nprose"


def test_a_renamed_doc_is_refused_with_a_useful_message(make_project):
    make_project("p1")
    with pytest.raises(ValueError) as err:
        import_service.import_bundle("p1", None, [("Old.docx", b"\xd0\xcf\x11\xe0not a zip")])
    assert ".docx" in str(err.value)


def test_a_word_document_mixed_with_markdown_imports_both(make_project):
    def fill(doc):
        doc.add_paragraph("Word prose.")

    make_project("p1")
    import_service.import_bundle(
        "p1",
        None,
        [("Bundle/01-word.docx", build_docx(fill)), ("Bundle/02-note.md", b"Markdown prose.")],
        {"name": "Bundle"},
    )
    titles = {doc["title"]: doc["body"] for doc in documents_service.iter_documents("p1")}
    assert titles == {"Word": "Word prose.", "Note": "Markdown prose."}


# ---------------------------------------------------------------------------
# round trip
# ---------------------------------------------------------------------------


def test_a_docx_exported_by_this_app_imports_back(make_project):
    """The strongest guarantee that the two sides agree: export, then import."""
    from app.services.projects import export_docx

    make_project("p1")
    documents_service.create_document(
        "p1", "Opening", kind="chapter", folder="Act One", content="It began quietly.\n\nThen it did not."
    )
    documents_service.create_document(
        "p1", "Second Scene", kind="chapter", folder="Act One", content="**Bold** and *italic* text."
    )

    raw = export_docx("p1")
    make_project("p2")
    import_service.import_bundle("p2", None, [("round-trip.docx", raw)], {"split": True})

    imported = documents_service.iter_documents("p2")
    by_title = {doc["title"]: doc["body"] for doc in imported}
    assert "Opening" in by_title, f"expected the exported chapter titles, got {list(by_title)}"
    assert "Second Scene" in by_title
    assert "It began quietly." in by_title["Opening"]
    assert "**Bold**" in by_title["Second Scene"]
