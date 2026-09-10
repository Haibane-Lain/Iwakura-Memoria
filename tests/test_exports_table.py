"""Exports must carry character tables (the wiki info box).

The box is raw HTML in the Markdown (``<aside class="character-table">`` with a
nested ``<table>``), so every exporter meets it: DOCX needs real Word-table
support (before this it silently dropped the rows), PDF renders the nested table
through fpdf2, and EPUB keeps the markup and embeds the portrait. A malformed or
empty box must never take an export down.
"""
from __future__ import annotations

import io
import re
import zipfile

from docx import Document
from PIL import Image

from app import config
from app.services import assets as assets_service
from app.services import documents as documents_service
from app.services import projects as projects_service


def png_bytes(width: int = 8, height: int = 6) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buf, "PNG")
    return buf.getvalue()


def _stored_image(name: str = "ubel.png") -> str:
    return assets_service.save_image("p1", name, png_bytes())["path"]


def _project_with(make_project, body: str = "") -> None:
    folder = make_project("p1")
    (folder / config.WIKI_DIRNAME).mkdir(exist_ok=True)
    documents_service.create_document("p1", "Chapter One", "chapter", content=body)


def _box(body: str = "Female", portrait: str | None = None) -> str:
    lines = ['<aside class="character-table" data-width="340">', '<table class="ct-rows">']
    lines.append('<tr class="ct-title"><th colspan="2">Übel</th></tr>')
    lines.append('<tr class="ct-subtitle"><th colspan="2">Anime • Manga</th></tr>')
    if portrait:
        lines.append(
            f'<tr class="ct-portrait"><td colspan="2"><p><img src="{portrait}" '
            'alt="Übel" width="240"></p></td></tr>'
        )
    lines.append('<tr class="ct-section"><th colspan="2">Biographical Information</th></tr>')
    lines.append(f'<tr class="ct-row"><td class="ct-label">Gender</td><td class="ct-value">{body}</td></tr>')
    lines.append('<tr class="ct-row"><td class="ct-label">Rank</td><td class="ct-value">Third-Class</td></tr>')
    lines.append("</table>")
    lines.append("</aside>")
    return "\n".join(lines)


def _docx(data: bytes) -> tuple[Document, zipfile.ZipFile]:
    archive = zipfile.ZipFile(io.BytesIO(data))
    return Document(io.BytesIO(data)), archive


def _table_text(doc: Document) -> list[list[str]]:
    return [
        [cell.text.strip() for cell in row.cells]
        for table in doc.tables
        for row in table.rows
    ]


def _pdf_text(data: bytes) -> bytes:
    """The text of a PDF, with fpdf2's compressed content streams inflated.

    Only meaningful when the export uses fpdf2's built-in core fonts: a Unicode
    TTF is embedded as a subset with its own encoding, so the words are no
    longer literal bytes in the stream (see ``_core_fonts``).
    """
    import zlib

    text = b""
    for match in re.finditer(rb"stream\r?\n", data):
        start = match.end()
        end = data.find(b"endstream", start)
        chunk = data[start:end]
        try:
            text += zlib.decompress(chunk)
        except zlib.error:
            text += chunk
    return text


def _core_fonts(monkeypatch) -> None:
    """Force fpdf2's built-in fonts so the PDF's words can be asserted."""
    monkeypatch.setattr(
        projects_service, "_find_pdf_fonts", lambda: {"serif": None, "mono": None}
    )


def _epub_parts(data: bytes) -> tuple[zipfile.ZipFile, str]:
    archive = zipfile.ZipFile(io.BytesIO(data))
    html = "\n".join(
        archive.read(name).decode("utf-8", "replace")
        for name in archive.namelist()
        if name.endswith(".xhtml")
    )
    return archive, html


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def test_docx_keeps_the_box_as_a_real_table(make_project):
    _project_with(make_project, _box())

    doc, _ = _docx(projects_service.export_docx("p1"))
    assert len(doc.tables) == 1
    rows = _table_text(doc)
    assert [row[0] for row in rows] == [
        "Übel",
        "Anime • Manga",
        "Biographical Information",
        "Gender",
        "Rank",
    ]
    # The label/value pair keeps its two cells.
    assert rows[3][1] == "Female"
    assert rows[4][1] == "Third-Class"


def test_docx_embeds_the_portrait_inside_the_table(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document("p1", "01-chapter-one", _box(portrait=path))

    doc, archive = _docx(projects_service.export_docx("p1"))
    assert len(doc.tables) == 1
    media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert media, "the portrait is embedded"
    assert "Übel" in _table_text(doc)[0][0]


def test_docx_survives_a_malformed_box(make_project):
    _project_with(make_project)
    body = (
        '<aside class="character-table">\n<table class="ct-rows">\n'
        '<tr class="ct-title"><th colspan="2">No rows here</th></tr>\n'
        '<tr class="ct-row"><td class="ct-label">Only a label</td></tr>\n'
        '<tr class="ct-row"><td class="ct-label">Missing</td>'
        '<td class="ct-value"><img src="assets/gone.png"></td></tr>\n'
        "</table>\n</aside>"
    )
    documents_service.save_document("p1", "01-chapter-one", body)

    doc, _ = _docx(projects_service.export_docx("p1"))
    assert len(doc.tables) == 1
    assert _table_text(doc)[1][0] == "Only a label"


def test_docx_handles_a_document_without_a_box(make_project):
    _project_with(make_project, "Just prose.\n\n- and a list")

    doc, _ = _docx(projects_service.export_docx("p1"))
    assert doc.tables == []
    assert any("Just prose." in p.text for p in doc.paragraphs)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_pdf_renders_the_box_table(make_project, monkeypatch):
    _project_with(make_project, _box())
    _core_fonts(monkeypatch)

    data = projects_service.export_pdf("p1")
    assert data[:5] == b"%PDF-"
    text = _pdf_text(data)
    assert "Übel".encode("latin-1") in text
    assert b"Gender" in text
    assert b"Female" in text
    assert b"Biographical Information" in text
    # Two columns: the label and its value sit at different x offsets on the
    # same line of the table.
    offsets = [float(value) for value in re.findall(rb"([\d.]+) [\d.]+ Td", text)]
    assert len(set(round(value) for value in offsets)) > 1


def test_pdf_embeds_a_portrait_inside_the_box(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document("p1", "01-chapter-one", _box(portrait=path))

    data = projects_service.export_pdf("p1")
    assert data[:5] == b"%PDF-"
    assert b"/Subtype /Image" in data


def test_pdf_survives_a_box_with_a_missing_portrait(make_project, monkeypatch):
    _project_with(make_project, _box(portrait="assets/gone.png"))
    _core_fonts(monkeypatch)

    data = projects_service.export_pdf("p1")
    assert data[:5] == b"%PDF-"
    assert b"/Subtype /Image" not in data
    assert b"Female" in _pdf_text(data)


def test_pdf_core_font_fallback_does_not_raise_on_a_bullet(make_project, monkeypatch):
    # The reference info box's subtitle is "Anime • Manga • Young": with the
    # latin-1 core fonts that character used to raise inside fpdf2 and take the
    # export down.
    _project_with(make_project, _box())
    _core_fonts(monkeypatch)

    data = projects_service.export_pdf("p1")
    assert data[:5] == b"%PDF-"
    text = _pdf_text(data)
    assert b"?" in text
    assert b"Female" in text


# ---------------------------------------------------------------------------
# EPUB
# ---------------------------------------------------------------------------


def test_epub_keeps_the_box_markup_and_styles_it(make_project):
    _project_with(make_project, _box())

    archive, html = _epub_parts(projects_service.export_epub("p1"))
    assert '<aside class="character-table"' in html
    assert '<table class="ct-rows">' in html
    assert '<td class="ct-label">Gender</td>' in html
    styles = [
        archive.read(name).decode("utf-8", "replace")
        for name in archive.namelist()
        if name.endswith(".css")
    ]
    assert any("aside.character-table" in sheet for sheet in styles), styles


def test_epub_embeds_the_portrait_of_a_box(make_project):
    _project_with(make_project)
    path = _stored_image()
    documents_service.save_document("p1", "01-chapter-one", _box(portrait=path))

    archive, html = _epub_parts(projects_service.export_epub("p1"))
    images = [
        name[len("EPUB/"):] if name.startswith("EPUB/") else name
        for name in archive.namelist()
    ]
    portraits = [name for name in images if name.startswith("images/")]
    assert len(portraits) == 1
    assert f'src="{portraits[0]}"' in html
    assert "assets/" not in html
